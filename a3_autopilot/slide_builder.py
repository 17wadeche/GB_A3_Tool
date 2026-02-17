from __future__ import annotations
from datetime import date
from pathlib import Path
from typing import Iterable
from pptx import Presentation
from pptx.exc import PackageNotFoundError
from a3_autopilot.models import DmaicPackage
def _iter_shapes_recursive(shape):
    yield shape
    if hasattr(shape, "shapes"):  # group shape
        for sub in shape.shapes:
            yield from _iter_shapes_recursive(sub)
def _replace_in_slide(slide, replacements: dict[str, str]) -> None:
    for shape in slide.shapes:
        for sub in _iter_shapes_recursive(shape):
            if not hasattr(sub, "text"):
                continue
            original = (sub.text or "").strip()
            if not original:
                continue
            for marker, new_text in replacements.items():
                if marker.lower() in original.lower():
                    sub.text = new_text
                    break
def _build_section_text(pkg: DmaicPackage) -> dict[str, str]:
    ms = pkg.measure_summary
    goal = (
        f"{pkg.define.goal_metric.metric_name}: "
        f"{pkg.define.goal_metric.baseline if pkg.define.goal_metric.baseline is not None else 'current'} "
        f"-> {pkg.define.goal_metric.target} by {pkg.define.goal_metric.due_date}"
    )
    top_roots = [f"- {rc.id} [{rc.category}] {rc.statement[:110]}" for rc in pkg.root_causes[:3]]
    top_cm = [
        f"- {cm.id} ({'PRIMARY' if cm.is_primary else 'BACKUP' if cm.is_backup else 'ALT'}) {cm.description[:95]}"
        for cm in pkg.countermeasures[:3]
    ]
    actions = [f"- {a.owner} | due {a.due_date} | KPI {a.kpi}" for a in pkg.actions[:3]]
    define_text = (
        "DEFINE: Describe the Problem\n"
        f"Problem Statement:\n{pkg.define.problem_statement}\n\n"
        f"Business Impact:\n{pkg.define.business_impact or 'Not provided'}\n\n"
        f"Scope In: {pkg.define.scope_in}\n"
        f"Scope Out: {pkg.define.scope_out}\n\n"
        f"Goal:\n{goal}\n"
        f"Team: {', '.join([f'{m.name} ({m.role})' for m in pkg.define.team])}"
    )
    measure_text = (
        "MEASURE: Identify the Current State\n"
        f"Current Performance:\n"
        f"- Baseline Mean Impact: {ms.get('baseline_mean_impact', 'N/A')}\n"
        f"- Total Impact: {ms.get('baseline_total_impact', 'N/A')}\n"
        f"- Records: {ms.get('records', 0)}\n"
        f"- Vital Few Categories: {ms.get('vital_few_count', 0)}\n\n"
        "Process Thinking:\n"
        f"- Pareto built from mapped fields\n"
        f"- Data confidence: {pkg.confidence_score:.2f}\n"
        f"- Notes: {'; '.join(pkg.confidence_notes) if pkg.confidence_notes else 'No major data gaps'}"
    )
    analyze_text = (
        "ANALYZE: Find the Root Cause(s)\n"
        "Validated Root Causes:\n"
        + ("\n".join(top_roots) if top_roots else "- None identified")
        + "\n\nFive Whys Summary:\n"
        + "\n".join([f"- Why {w.level}: {w.why[:95]} (conf {w.confidence:.2f})" for w in pkg.five_whys[:5]])
    )
    improve_text = (
        "IMPROVE: Optimize and Act\n"
        "Countermeasures:\n"
        + ("\n".join(top_cm) if top_cm else "- None proposed")
        + "\n\nActions:\n"
        + ("\n".join(actions) if actions else "- No actions generated")
    )
    cp = pkg.control_plan
    control_text = (
        "CONTROL: Demonstrate Improvement & Sustainability\n"
        f"Cadence: {cp.cadence}\n"
        f"Audit Frequency: {cp.audit_frequency}\n"
        f"Response Plan: {cp.response_plan}\n"
        f"Standard Work: {cp.standard_work_update}\n\n"
        f"Next Review: {date.today()}\n"
        f"Assumptions: {' | '.join([a.text for a in pkg.assumptions]) if pkg.assumptions else 'None'}"
    )
    return {
        "define: describe the problem": define_text,
        "measure: identify the current state": measure_text,
        "analyze: find the root cause": analyze_text,
        "improve: optimize and act": improve_text,
        "control: demonstrate improvement": control_text,
    }
def _candidate_template_paths(explicit: str | Path | None = None) -> Iterable[Path]:
    if explicit:
        yield Path(explicit)
    yield Path("templates") / "A3 Template.pptx"
    yield Path("examples") / "A3 Template.pptx"
    yield Path("A3 Template.pptx")
    mod_root = Path(__file__).resolve().parents[1]
    yield mod_root / "templates" / "A3 Template.pptx"
    yield mod_root / "examples" / "A3 Template.pptx"
    yield Path("/mnt/data/A3 Template.pptx")
def _load_template_presentation(template_path: str | Path | None) -> Presentation:
    last_err = None
    for p in _candidate_template_paths(template_path):
        try:
            if p.exists():
                return Presentation(str(p))
        except (PackageNotFoundError, OSError) as e:
            last_err = e
            continue
    msg = (
        "A3 template PPTX not found. Place 'A3 Template.pptx' in one of:\n"
        " - ./templates/\n - ./examples/\n - project root\n"
        "or pass template_path explicitly."
    )
    if last_err:
        msg += f"\nLast error: {last_err}"
    raise FileNotFoundError(msg)
def render_single_slide(
    pkg: DmaicPackage,
    output_path: str | Path,
    template_path: str | Path | None = None,
) -> str:
    prs = _load_template_presentation(template_path)
    slide = prs.slides[0]
    replacements = _build_section_text(pkg)
    _replace_in_slide(slide, replacements)
    output_path = str(output_path)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path